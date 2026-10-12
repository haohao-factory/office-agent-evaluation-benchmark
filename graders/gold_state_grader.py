import json, hashlib
from pathlib import Path
from openpyxl import load_workbook
from docx import Document
def _hash(p):
 h=hashlib.sha256()
 for x in sorted(y for y in p.rglob('*') if y.is_file()): h.update(str(x.relative_to(p)).encode());h.update(x.read_bytes())
 return h.hexdigest()
def grade(root,run_id,task_id):
 root=Path(root);spec=json.loads((root/'environment/task_state_mapping/smoke_gold_state_v0.1.json').read_text(encoding='utf-8-sig'))['tasks'][task_id];base=root/'environment/base_state';state=root/'environment/runtime/runs'/run_id
 if spec.get('no_side_effect'): return {'state_status':'PASS' if _hash(base)==_hash(state) else 'FAIL','checks':['no_side_effect']}
 if 'path_exists' in spec:return {'state_status':'PASS' if (state/spec['path_exists']).exists() and not (state/spec['path_absent']).exists() else 'FAIL','checks':['path_change']}
 if spec['entity'].endswith('.xlsx'):
  ws=load_workbook(state/spec['entity'])[spec['sheet']];h=[c.value for c in ws[1]];row=next((r for r in ws.iter_rows(min_row=2,values_only=True) if all(r[h.index(k)]==v for k,v in spec['match'].items())),None);ok=row is not None and all(row[h.index(k)]==v for k,v in spec['equals'].items());return {'state_status':'PASS' if ok else 'FAIL','checks':['sheet_fields']}
 if spec['entity'].endswith('.docx'):
  d=Document(state/spec['entity']);text='\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells);return {'state_status':'PASS' if all(x in text for x in spec['contains']) else 'FAIL','checks':['doc_content']}
 if spec['entity'].endswith('.json'):
  all_state=json.loads((state/spec['entity']).read_text(encoding='utf-8'));xs=[x for x in all_state if all(x.get(k)==v for k,v in spec.get('match',{}).items())];ok=(not spec.get('count') or len(xs)==spec['count']) and bool(xs)
  if ok and spec.get('unchanged'):
   before=json.loads((base/spec['entity']).read_text(encoding='utf-8'));b=next((x for x in before if x.get('event_id')==xs[0].get('event_id')),{});ok=all(xs[0].get(k)==b.get(k) for k in spec['unchanged'])
  return {'state_status':'PASS' if ok else 'FAIL','checks':['json_object','must_not_change']}
 return {'state_status':'REVIEW','checks':['unsupported']}
